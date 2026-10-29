from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    UploadFile,
)
from sqlalchemy.orm import Session

from auth_utils import get_current_user

router = APIRouter()

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


@router.post("/extract")
async def extract_document(
    file: UploadFile = File(...),
    user=Depends(get_current_user),
):
    filename = file.filename or ""
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Supported formats: PDF, DOCX, and TXT.",
        )

    contents = await file.read(MAX_FILE_SIZE + 1)

    if not contents:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if len(contents) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File exceeds the 10 MB limit.",
        )

    try:
        if extension == ".txt":
            text = contents.decode("utf-8-sig")

        elif extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(contents))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDFs are not supported.",
                )

            text = "\n\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        else:
            from docx import Document

            document = Document(BytesIO(contents))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            # Include text from tables as well.
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(paragraphs)

    except HTTPException:
        raise
    except Exception as exc:
        print("Document extraction error:", repr(exc))
        raise HTTPException(
            status_code=400,
            detail="Could not extract text. The file may be invalid or corrupted.",
        )

    text = text.strip()

    if not text:
        raise HTTPException(
            status_code=422,
            detail=(
                "No text was found. If this is a scanned PDF, "
                "OCR is needed to recognize its text."
            ),
        )

    return {
        "filename": filename,
        "text": text,
        "characters": len(text),
    }
